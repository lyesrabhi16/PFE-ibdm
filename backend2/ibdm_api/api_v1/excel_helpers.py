from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def generateExcelReport(
    data,
    filename,
    title

):

    # Create a new workbook and select the active worksheet
    wb = Workbook()
    ws = wb.active
    if not ws:
        raise ValueError("No active worksheet found!") 
    # Set the title of the worksheet
    ws.title = title
    ws.append(["Rapport", data["rapport_name"]])
    ws.append(["Interval", data["interval"]])
    ws.append(["Date", data["date"]])
    ws.append([])  # Add an empty row for spacing
    
    
    # Write the header row
    header_row = [header['title'] for header in data['headers']]
    ws.append(header_row)

    # Apply styling to the header row
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="009879", end_color="009879", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    for col in range(1, len(header_row) + 1):
        cell = ws.cell(row=5, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    # Write the data rows
    for row in data['rows']:
        row_data = []
        for code in data['header_codes']:
            value = row.get(code, '')
            if code in data['suffixes']:
                value = f"{value}{data['suffixes'][code]}"
            row_data.append(value)
        ws.append(row_data)

        # Apply conditional formatting
        row_index = ws.max_row
        if row == data["rows"][-1]:
            for col in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_index, column=col)
                cell.fill = PatternFill(start_color="009879", end_color="009879", fill_type="solid")
                cell.font = Font(bold=True, color="FFFFFF")
        elif row.get('highlight'):
            for col in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_index, column=col)
                cell.fill = PatternFill(start_color="F3F3F3", end_color="F3F3F3", fill_type="solid")
                cell.font = Font(bold=True, color="23db0b")
        if row.get('danger'):
            for col in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_index, column=col)
                cell.font = Font(color="FF0000")

    # Write the synthese section
    ws.append([])  # Add an empty row for spacing
    ws.append(["Synthèse"])
    ws.append([data['synthese']])

    # Apply styling to the synthese section
    synthese_font = Font(bold=True, color="009879")
    ws.cell(row=ws.max_row - 1, column=1).font = synthese_font

    # Adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter  # Get the column name
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = (max_length + 2) * 1.2
        ws.column_dimensions[column].width = adjusted_width

    # Save the workbook
    output_file = f"{filename}.xlsx"
    wb.save(output_file)

    print(f"Excel file '{output_file}' created successfully!")

    return output_file